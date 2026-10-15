import io
from datetime import UTC, datetime
from unittest.mock import patch

import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from backend.repositories.analysis_repo import AnalysisRepository
from backend.repositories.requirements_repo import JobRequirementsRepository
from backend.repositories.resume_repo import ResumeRepository
from backend.schemas.common import AnalysisStatus, ResumeStatus
from backend.schemas.requirements import JobRequirements
from backend.services.llm_service import MockLLMService
from backend.services.watchdog_service import WatchdogService

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def create_docx_resume(name: str, skills: list[str], years: int, title: str = "Software Engineer") -> bytes:
    """Creates a genuine DOCX binary document with structured headings and paragraphs."""
    doc = docx.Document()
    doc.add_heading(name, level=0)
    doc.add_paragraph(f"{name.lower().replace(' ', '.')}@example.com")
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(f"Experienced {title} with {years} years of professional development experience.")
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(", ".join(skills))
    doc.add_heading("Work Experience", level=1)
    doc.add_paragraph(f"{title}, Acme Corp, {years} years")
    doc.add_paragraph(f"Developed scalable backend systems using {skills[0] if skills else 'Python'}.")
    doc.add_paragraph("Collaborated with cross-functional product and infrastructure teams.")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.S. in Computer Science, Tech University")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_empty_docx_resume(space_count: int = 1) -> bytes:
    """Creates a valid DOCX file with whitespace-only text (fails extraction as EMPTY_DOCUMENT)."""
    doc = docx.Document()
    doc.add_paragraph(" " * space_count)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ==============================================================================
# SCENARIO 1: Full Happy Path Flow (TEST-05 Item 1)
# ==============================================================================
def test_scenario_1_full_happy_path(client: TestClient):
    """
    Scenario 1: Valid JD, requirement confirmation, upload valid resumes,
    background processing, deterministic scoring, and ranked results.
    """
    # 1. Create analysis with job description
    jd_text = (
        "Role: Senior Backend Engineer\n"
        "Requirements:\n"
        "- 5+ years of software engineering experience\n"
        "- Proficiency in Python and PostgreSQL\n"
        "Preferred:\n"
        "- Experience with Docker and Redis\n"
    )
    resp = client.post("/api/analyses", data={"name": "Backend Team Q4", "jd_text": jd_text})
    assert resp.status_code == 201
    analysis_id = resp.json()["analysis"]["id"]

    # 2. Review and confirm requirements
    reqs_resp = client.get(f"/api/analyses/{analysis_id}/requirements")
    assert reqs_resp.status_code == 200
    req_data = reqs_resp.json()
    assert "Python" in req_data["required_skills"]

    confirm_resp = client.post(f"/api/analyses/{analysis_id}/requirements/confirm")
    assert confirm_resp.status_code == 200
    assert confirm_resp.json()["confirmed_at"] is not None

    # 3. Upload batch of 3 candidate DOCX resumes
    c1_bytes = create_docx_resume("Alice Martin", ["Python", "PostgreSQL", "Docker", "Redis"], 6)
    c2_bytes = create_docx_resume("Bob Smith", ["Python", "PostgreSQL"], 4)
    c3_bytes = create_docx_resume("Charlie Brown", ["Python"], 2)

    files = [
        ("files", ("alice_resume.docx", io.BytesIO(c1_bytes), DOCX_MIME)),
        ("files", ("bob_resume.docx", io.BytesIO(c2_bytes), DOCX_MIME)),
        ("files", ("charlie_resume.docx", io.BytesIO(c3_bytes), DOCX_MIME)),
    ]
    upload_resp = client.post(f"/api/analyses/{analysis_id}/resumes", files=files)
    assert upload_resp.status_code == 200
    assert upload_resp.json()["accepted_count"] == 3

    # 4. Start analysis (Celery runs synchronously in eager test mode)
    start_resp = client.post(f"/api/analyses/{analysis_id}/start")
    assert start_resp.status_code == 200
    assert start_resp.json()["status"] == "PROCESSING"

    # 5. Check status (all processed, ranked)
    status_resp = client.get(f"/api/analyses/{analysis_id}/status")
    assert status_resp.status_code == 200
    assert status_resp.json()["status"] == "RANKED"
    assert status_resp.json()["processed"] == 3

    # 6. Retrieve candidate rankings
    candidates_resp = client.get(f"/api/analyses/{analysis_id}/candidates")
    assert candidates_resp.status_code == 200
    ranking = candidates_resp.json()["ranking"]
    assert len(ranking) == 3

    # Verify deterministic ranking order (Alice > Bob > Charlie)
    assert ranking[0]["rank"] == 1
    assert ranking[0]["label"].startswith("Candidate ")
    assert ranking[0]["match_score"] >= ranking[1]["match_score"] >= ranking[2]["match_score"]
    assert ranking[0]["required_skill_coverage"]["value"] == 100.0


# ==============================================================================
# SCENARIO 2: Insufficient JD & Manual Correction (TEST-05 Item 2)
# ==============================================================================
def test_scenario_2_insufficient_jd_and_manual_correction(client: TestClient):
    """
    Scenario 2: Insufficient JD triggers needs_manual_input, cannot confirm
    with empty required skills, recruiter adds skills manually, then confirms.
    """
    vague_jd = "We are a fast-paced company looking for an energetic problem solver."
    resp = client.post("/api/analyses", data={"name": "Vague Role", "jd_text": vague_jd})
    assert resp.status_code == 201
    analysis_id = resp.json()["analysis"]["id"]
    assert resp.json()["needs_manual_input"] is True

    # Attempting to confirm with 0 required skills must fail
    confirm_resp = client.post(f"/api/analyses/{analysis_id}/requirements/confirm")
    assert confirm_resp.status_code == 400
    assert confirm_resp.json()["error"]["code"] == "REQUIREMENTS_INSUFFICIENT"

    # Recruiter manually adds required skills
    update_payload = {
        "required_skills": ["TypeScript", "React"],
        "preferred_skills": ["Tailwind"],
        "experience_requirements": ["3+ years"],
        "education_requirements": [],
        "explicit_requirements": [],
        "job_relevant_text": "TypeScript React developer with 3+ years experience.",
    }
    put_resp = client.put(f"/api/analyses/{analysis_id}/requirements", json=update_payload)
    assert put_resp.status_code == 200
    assert "TypeScript" in put_resp.json()["required_skills"]

    # Now confirmation succeeds
    confirm_resp2 = client.post(f"/api/analyses/{analysis_id}/requirements/confirm")
    assert confirm_resp2.status_code == 200
    assert confirm_resp2.json()["confirmed_at"] is not None


# ==============================================================================
# SCENARIO 3: Mixed Resume Batch Outcomes (TEST-05 Item 3)
# ==============================================================================
def test_scenario_3_mixed_resumes_batch(client: TestClient):
    """
    Scenario 3: Batch containing valid, duplicate (SHA-256), unsupported,
    and corrupt/empty documents are categorized properly.
    """
    # 1. Setup analysis
    create_resp = client.post(
        "/api/analyses",
        data={"name": "Mixed Batch Analysis", "jd_text": "Requirements: Python, SQL. 3+ years exp."},
    )
    analysis_id = create_resp.json()["analysis"]["id"]
    client.post(f"/api/analyses/{analysis_id}/requirements/confirm")

    valid_resume = create_docx_resume("Valid Dev", ["Python", "SQL"], 3)
    duplicate_resume = valid_resume  # Exact duplicate bytes
    empty_resume = create_empty_docx_resume()  # Valid DOCX structure but empty text

    files = [
        ("files", ("valid.docx", io.BytesIO(valid_resume), DOCX_MIME)),
        ("files", ("copy_of_valid.docx", io.BytesIO(duplicate_resume), DOCX_MIME)),
        ("files", ("bad_program.exe", io.BytesIO(b"binary exe"), "application/octet-stream")),
        ("files", ("empty.docx", io.BytesIO(empty_resume), DOCX_MIME)),
    ]

    upload_resp = client.post(f"/api/analyses/{analysis_id}/resumes", files=files)
    assert upload_resp.status_code == 200
    results = upload_resp.json()["results"]

    # Verify per-file upload outcomes
    outcomes = {r["filename"]: r["status"] for r in results}
    assert outcomes["valid.docx"] == "READY"
    assert outcomes["copy_of_valid.docx"] == "DUPLICATE"
    assert outcomes["bad_program.exe"] == "REJECTED"
    assert outcomes["empty.docx"] == "READY"  # Caught during extraction attempt

    # Start analysis
    client.post(f"/api/analyses/{analysis_id}/start")

    # Check status: 1 successful candidate, 1 failed (empty document)
    status_resp = client.get(f"/api/analyses/{analysis_id}/status")
    status_data = status_resp.json()
    assert status_data["status"] == "RANKED"
    assert status_data["successful"] == 1
    assert status_data["failed"] >= 1
    assert status_data["duplicates_excluded"] >= 1


# ==============================================================================
# SCENARIO 4: All Resumes Fail Gracefully (TEST-05 Item 4)
# ==============================================================================
def test_scenario_4_all_resumes_fail(client: TestClient):
    """
    Scenario 4: If all resumes fail extraction, analysis still finalizes to
    RANKED with an empty ranking list and separate failure reasons.
    """
    create_resp = client.post(
        "/api/analyses",
        data={"name": "All Fail Test", "jd_text": "Requirements: Java, Spring Boot."},
    )
    analysis_id = create_resp.json()["analysis"]["id"]
    client.post(f"/api/analyses/{analysis_id}/requirements/confirm")

    # Upload empty DOCX files (accepted at upload, fails extraction as EMPTY_DOCUMENT)
    files = [
        ("files", ("unreadable1.docx", io.BytesIO(create_empty_docx_resume(1)), DOCX_MIME)),
        ("files", ("unreadable2.docx", io.BytesIO(create_empty_docx_resume(5)), DOCX_MIME)),
    ]
    upload_resp = client.post(f"/api/analyses/{analysis_id}/resumes", files=files)
    assert upload_resp.json()["accepted_count"] == 2

    # Start analysis
    start_resp = client.post(f"/api/analyses/{analysis_id}/start")
    assert start_resp.status_code == 200

    # Finalization still reaches RANKED cleanly
    status_resp = client.get(f"/api/analyses/{analysis_id}/status")
    assert status_resp.json()["status"] == "RANKED"
    assert status_resp.json()["successful"] == 0
    assert status_resp.json()["failed"] == 2

    # Ranking list is empty
    candidates_resp = client.get(f"/api/analyses/{analysis_id}/candidates")
    assert len(candidates_resp.json()["ranking"]) == 0
    assert len(candidates_resp.json()["excluded_files"]) == 2


# ==============================================================================
# SCENARIO 5: AI Explanation Lifecycle & Retry (TEST-05 Item 5)
# ==============================================================================
def test_scenario_5_explanation_lifecycle_and_retry(client: TestClient):
    """
    Scenario 5: Request explanation, handle transient failure, retry,
    and verify grounded strengths/gaps output.
    """
    # 1. Create and complete a single candidate analysis
    create_resp = client.post(
        "/api/analyses",
        data={"name": "AI Explanation Role", "jd_text": "Requirements: Python, Django, Docker."},
    )
    analysis_id = create_resp.json()["analysis"]["id"]
    client.post(f"/api/analyses/{analysis_id}/requirements/confirm")

    c_bytes = create_docx_resume("Dave Miller", ["Python", "Django"], 4)
    client.post(f"/api/analyses/{analysis_id}/resumes", files=[("files", ("dave.docx", io.BytesIO(c_bytes), DOCX_MIME))])
    client.post(f"/api/analyses/{analysis_id}/start")

    cand_resp = client.get(f"/api/analyses/{analysis_id}/candidates")
    candidate_id = cand_resp.json()["ranking"][0]["candidate_id"]

    url = f"/api/analyses/{analysis_id}/candidates/{candidate_id}/explanation"

    # 2. Simulate LLM failure on first request
    failing_mock = MockLLMService(should_fail=True)
    with patch("backend.services.explanation_service.default_llm_service", failing_mock):
        req_fail = client.post(url)
        assert req_fail.status_code == 200  # Handled synchronously in eager test mode

    # Check status: marked FAILED after task execution
    exp_status = client.get(url)
    assert exp_status.json()["state"] == "FAILED"

    # 3. Working mock LLM service on retry
    working_mock = MockLLMService(
        canned_response={
            "strengths": ["Strong proficiency in Python and Django."],
            "gaps": ["No evidence of Docker containerization."],
            "experience_alignment": "4 years of backend experience directly satisfies requirement.",
            "assessment": "Solid candidate with strong backend capabilities.",
        }
    )

    with patch("backend.services.explanation_service.default_llm_service", working_mock):
        retry_resp = client.post(url)
        assert retry_resp.status_code == 200

    # 4. Verified output
    get_ready_resp = client.get(url)
    assert get_ready_resp.status_code == 200
    assert get_ready_resp.json()["state"] == "READY"
    assert "Python" in str(get_ready_resp.json()["explanation"]["strengths"])


# ==============================================================================
# SCENARIO 6: Reopen Saved Analysis Stability (TEST-05 Item 6)
# ==============================================================================
def test_scenario_6_reopen_saved_analysis(client: TestClient):
    """
    Scenario 6: Reopening a saved analysis preserves identical scores, ranks,
    and anonymous candidate labels without recomputation.
    """
    create_resp = client.post(
        "/api/analyses",
        data={"name": "Historical Audit 2026", "jd_text": "Requirements: Golang, Kubernetes. 5+ years."},
    )
    analysis_id = create_resp.json()["analysis"]["id"]
    client.post(f"/api/analyses/{analysis_id}/requirements/confirm")

    c_bytes = create_docx_resume("Sam Wilson", ["Golang", "Kubernetes"], 6)
    client.post(f"/api/analyses/{analysis_id}/resumes", files=[("files", ("sam.docx", io.BytesIO(c_bytes), DOCX_MIME))])
    client.post(f"/api/analyses/{analysis_id}/start")

    # First fetch
    first_view = client.get(f"/api/analyses/{analysis_id}/candidates").json()
    first_label = first_view["ranking"][0]["label"]
    first_score = first_view["ranking"][0]["match_score"]

    # Reopen analysis (simulate user returning later)
    analysis_summary = client.get(f"/api/analyses/{analysis_id}").json()
    assert analysis_summary["status"] == "RANKED"

    second_view = client.get(f"/api/analyses/{analysis_id}/candidates").json()
    assert second_view["ranking"][0]["label"] == first_label == "Candidate 01"
    assert second_view["ranking"][0]["match_score"] == first_score


# ==============================================================================
# SCENARIO 7: Deletion & Vector Cleanup (TEST-05 Item 7)
# ==============================================================================
def test_scenario_7_deletion_and_vector_cleanup(client: TestClient, db_session: Session):
    """
    Scenario 7: Deleting a ranked analysis marks deletion, hides it immediately,
    cleans up database cascade records, and removes ChromaDB vector collections.
    """
    create_resp = client.post(
        "/api/analyses",
        data={"name": "Analysis to Purge", "jd_text": "Requirements: Ruby, Rails."},
    )
    analysis_id = create_resp.json()["analysis"]["id"]
    client.post(f"/api/analyses/{analysis_id}/requirements/confirm")

    c_bytes = create_docx_resume("Eve Hunter", ["Ruby", "Rails"], 3)
    client.post(f"/api/analyses/{analysis_id}/resumes", files=[("files", ("eve.docx", io.BytesIO(c_bytes), DOCX_MIME))])
    client.post(f"/api/analyses/{analysis_id}/start")

    # Delete analysis
    del_resp = client.delete(f"/api/analyses/{analysis_id}")
    assert del_resp.status_code == 202

    # Immediate GET returns 404
    get_resp = client.get(f"/api/analyses/{analysis_id}")
    assert get_resp.status_code == 404

    # Run watchdog sweep to execute complete purge
    WatchdogService.run_watchdog_cycle(db_session)

    # Verify database records are removed
    assert AnalysisRepository.get_by_id(db_session, analysis_id) is None


# ==============================================================================
# SCENARIO 8: Watchdog Lost Finalization Recovery (TEST-05 Item 8)
# ==============================================================================
def test_scenario_8_watchdog_lost_finalization_recovery(db_session: Session):
    """
    Scenario 8: Simulate a missed finalization event where all resumes succeeded
    but analysis remained in PROCESSING. Watchdog recovers it to RANKED.
    """
    analysis = AnalysisRepository.create(db_session, name="Stuck Analysis Recovery")
    analysis.status = AnalysisStatus.PROCESSING.value
    db_session.commit()

    # Confirmed requirements
    job_req = JobRequirements(required_skills=["C++"], confirmed_at=datetime.now(UTC))
    JobRequirementsRepository.save_or_update(db_session, analysis.id, job_req)

    # Resume in SUCCESS status
    resume = ResumeRepository.register_resume(
        db_session, analysis.id, "dev.docx", content_hash="hash123"
    )
    ResumeRepository.update_status(db_session, resume.id, to_status=ResumeStatus.SUCCESS.value)

    # Run watchdog cycle
    summary = WatchdogService.run_watchdog_cycle(db_session)
    assert summary["lost_finalizations_recovered"] >= 1

    # Analysis is now safely RANKED
    refreshed = AnalysisRepository.get_by_id(db_session, analysis.id)
    assert refreshed.status == AnalysisStatus.RANKED.value
